#!/usr/bin/env python3
"""DOCX 导入器的真实文件回归测试。"""

from __future__ import annotations

import json
import re
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path

import yaml
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
IMPORTER = ROOT / "importer" / "docx-importer.py"


def front_matter(markdown: str) -> dict:
    _, raw, _ = markdown.split("---", 2)
    return yaml.safe_load(raw)


class DocxImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="writing-docx-test-")
        self.root = Path(self.temp.name)
        (self.root / "content").mkdir()
        self.source = self.root / "《测试文集·春日记》.docx"

    def tearDown(self):
        self.temp.cleanup()

    def run_import(self, *extra: str, expected: int = 0):
        command = [
            sys.executable,
            str(IMPORTER),
            str(self.source),
            "--project-root",
            str(self.root),
            "--section",
            "works",
            "--title",
            "春日记",
            "--slug",
            "spring-notes",
            "--collection",
            "示例文集",
            "--collection-slug",
            "sample-works",
            "--category",
            "纪事",
            "--tag",
            "春日, 测试",
            "--json",
            *extra,
        ]
        completed = subprocess.run(command, text=True, encoding="utf-8", capture_output=True, timeout=45)
        self.assertEqual(completed.returncode, expected, completed.stderr)
        return json.loads(completed.stdout) if expected == 0 else completed

    def basic_document(self):
        doc = Document()
        doc.add_paragraph("春日记", style="Title")
        doc.add_heading("第一章", level=1)
        paragraph = doc.add_paragraph("这是一个保留中文标点，且不会把每句话拆开的自然段。")
        paragraph.add_run("粗体").bold = True
        paragraph.add_run("与")
        paragraph.add_run("斜体").italic = True
        doc.add_heading("第一节", level=2)
        doc.add_paragraph("史家记录的不只是结果。", style="Quote")
        doc.save(self.source)

    def test_01普通Word转换并保留自然段(self):
        self.basic_document()
        result = self.run_import()
        markdown = Path(result["markdown"]).read_text(encoding="utf-8")
        self.assertIn("这是一个保留中文标点，且不会把每句话拆开的自然段。", markdown)
        self.assertNotIn("春日记\n\n# 第一章", markdown)

    def test_02中文标题与多级标题(self):
        self.basic_document()
        result = self.run_import()
        markdown = Path(result["markdown"]).read_text(encoding="utf-8")
        self.assertEqual(front_matter(markdown)["title"], "春日记")
        self.assertIn("# 第一章", markdown)
        self.assertIn("## 第一节", markdown)

    def test_03粗体斜体与引用(self):
        self.basic_document()
        result = self.run_import()
        markdown = Path(result["markdown"]).read_text(encoding="utf-8")
        self.assertIn("**粗体**", markdown)
        self.assertIn("*斜体*", markdown)
        self.assertIn("> 史家记录的不只是结果。", markdown)

    def test_04_1Word居中和右对齐段落导入(self):
        doc = Document()
        doc.add_paragraph("山有木兮木有枝")
        centered = doc.add_paragraph("心悦君兮君不知")
        centered.alignment = WD_ALIGN_PARAGRAPH.CENTER
        signed = doc.add_paragraph("君纪鉴")
        signed.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        plain = doc.add_paragraph("普通左对齐正文")
        justified = doc.add_paragraph("两端对齐正文")
        justified.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        doc.save(self.source)
        result = self.run_import()
        markdown = Path(result["markdown"]).read_text(encoding="utf-8")
        self.assertIn(
            "{{< align center >}}\n\n心悦君兮君不知\n\n{{< /align >}}",
            markdown,
            "Word 居中段落应映射为 align center 短代码",
        )
        self.assertIn(
            "{{< align right >}}\n\n君纪鉴\n\n{{< /align >}}",
            markdown,
            "Word 右对齐段落应映射为 align right 短代码",
        )
        self.assertIn("普通左对齐正文", markdown)
        self.assertNotIn("{{< align left >}}", markdown, "左对齐不应产生冗余标记")
        self.assertNotIn("{{< align", markdown.split("两端对齐正文")[0][-30:], "JUSTIFY 应降级为左对齐")

    def test_04图片提取原图备份并转WebP(self):
        image = self.root / "source.png"
        Image.new("RGB", (1200, 800), "#9a4d3a").save(image)
        doc = Document()
        doc.add_paragraph("含图片的正文。")
        doc.add_picture(str(image))
        doc.save(self.source)
        result = self.run_import()
        bundle = Path(result["output"])
        markdown = (bundle / "index.md").read_text(encoding="utf-8")
        self.assertEqual(result["image_count"], 1)
        self.assertTrue((bundle / "images" / "image-001.webp").is_file())
        self.assertTrue((bundle / "images" / "original" / "image-001.png").is_file())
        self.assertIn("](images/image-001.webp)", markdown)

    def test_05FrontMatter与PageBundle路径(self):
        self.basic_document()
        result = self.run_import("--no-draft")
        bundle = self.root / "content" / "works" / "sample-works" / "spring-notes"
        self.assertEqual(Path(result["output"]).resolve(), bundle.resolve())
        data = front_matter((bundle / "index.md").read_text(encoding="utf-8"))
        self.assertEqual(data["slug"], "spring-notes")
        self.assertEqual(data["collections"], ["示例文集"])
        self.assertEqual(data["categories"], ["纪事"])
        self.assertEqual(data["tags"], ["春日", "测试"])
        self.assertFalse(data["draft"])
        self.assertTrue(data["description"])
        collection = self.root / "content" / "works" / "sample-works" / "_index.md"
        self.assertTrue(collection.is_file())
        self.assertEqual(front_matter(collection.read_text(encoding="utf-8"))["title"], "示例文集")

    def test_06特殊字符不会注入Markdown或HTML(self):
        doc = Document()
        doc.add_paragraph("<script>alert('x')</script> 与 *星号* [括号] #井号")
        doc.save(self.source)
        result = self.run_import()
        markdown = Path(result["markdown"]).read_text(encoding="utf-8")
        self.assertIn("&lt;script&gt;", markdown)
        body = markdown.split("---", 2)[2]
        self.assertNotIn("<script>", body)
        self.assertIn("\\*星号\\*", markdown)
        self.assertIn("\\[括号\\]", markdown)

    def test_07十万字长文可在合理时间导入(self):
        doc = Document()
        doc.add_heading("十万字压力测试", level=1)
        paragraph = "这是用于验证十万字Word导入能力的占位中文段落，不是真实作品。"
        cjk_per_paragraph = len(re.findall(r"[\u3400-\u4dbf\u4e00-\u9fff]", paragraph))
        repetitions = 100_000 // cjk_per_paragraph + 1
        for _ in range((repetitions + 79) // 80):
            doc.add_paragraph(paragraph * min(80, repetitions))
            repetitions -= min(80, repetitions)
            if repetitions <= 0:
                break
        doc.save(self.source)
        started = time.monotonic()
        result = self.run_import()
        elapsed = time.monotonic() - started
        self.assertGreaterEqual(result["word_count"], 95_000)
        self.assertLess(elapsed, 30)
        self.assertGreater(Path(result["markdown"]).stat().st_size, 200_000)

    def test_08拒绝覆盖已有文章且不破坏原文件(self):
        self.basic_document()
        first = self.run_import()
        markdown = Path(first["markdown"])
        before = markdown.read_bytes()
        second = self.run_import(expected=2)
        self.assertIn("目标已存在", second.stderr)
        self.assertEqual(markdown.read_bytes(), before)

    def test_09拒绝宏文件和非法slug(self):
        macro = self.root / "unsafe.docm"
        macro.write_bytes(b"not a document")
        command = [
            sys.executable,
            str(IMPORTER),
            str(macro),
            "--project-root",
            str(self.root),
            "--section",
            "essays",
            "--title",
            "测试",
            "--slug",
            "../unsafe",
        ]
        completed = subprocess.run(command, text=True, encoding="utf-8", capture_output=True)
        self.assertEqual(completed.returncode, 2)
        self.assertIn("只接受.docx", completed.stderr)

    def test_10外部链接不访问网络并产生提示(self):
        doc = Document()
        paragraph = doc.add_paragraph("请参看")
        part = paragraph.part
        relationship_id = part.relate_to(
            "https://example.invalid/never-requested",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
        hyperlink = paragraph._p.makeelement(qn("w:hyperlink"), {qn("r:id"): relationship_id})
        run = hyperlink.makeelement(qn("w:r"))
        text = run.makeelement(qn("w:t"))
        text.text = "外部资料"
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        doc.save(self.source)
        result = self.run_import()
        self.assertTrue(any(item["code"] == "external-link" for item in result["warnings"]))
        self.assertIn("外部资料", Path(result["markdown"]).read_text(encoding="utf-8"))

    def test_11重新导入前永久备份旧PageBundle(self):
        self.basic_document()
        first = self.run_import()
        old_markdown = Path(first["markdown"]).read_text(encoding="utf-8")
        doc = Document()
        doc.add_paragraph("这是重新导入后的新正文。")
        doc.save(self.source)
        backup_root = self.root / "author-backups"
        second = self.run_import("--force", "--backup-dir", str(backup_root))
        self.assertTrue(second["backup"])
        backup_markdown = Path(second["backup"]) / "index.md"
        self.assertEqual(backup_markdown.read_text(encoding="utf-8"), old_markdown)
        self.assertIn("新正文", Path(second["markdown"]).read_text(encoding="utf-8"))
if __name__ == "__main__":
    unittest.main(verbosity=2)
