#!/usr/bin/env python3
"""整部文集重新导入的元数据保持测试。"""

from __future__ import annotations

import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import yaml
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
IMPORTER = ROOT / "importer" / "docx-series-importer.py"


def source(data: dict, body: str = "旧正文。") -> str:
    return "---\n" + yaml.safe_dump(data, allow_unicode=True, sort_keys=False) + "---\n\n" + body + "\n"


def front_matter(path: Path) -> dict:
    _, raw, _ = path.read_text(encoding="utf-8").split("---", 2)
    return yaml.safe_load(raw)


class DocxSeriesImporterTests(unittest.TestCase):
    def test_重新导入保留首次发布日期稳定标识与自定义元数据(self):
        with tempfile.TemporaryDirectory(prefix="writing-series-test-") as directory:
            project = Path(directory)
            collection = project / "content" / "works" / "sample-collection"
            first = collection / "first"
            second = collection / "second"
            first.mkdir(parents=True)
            second.mkdir()
            (collection / "_index.md").write_text(
                source(
                    {
                        "title": "示例文集",
                        "date": "2024-01-01",
                        "lastmod": "2024-01-02",
                        "endDate": "2024-01-01",
                        "customCollectionField": "保留",
                    },
                    "",
                ),
                encoding="utf-8",
            )
            common = {
                "subtitle": "",
                "lastmod": "2024-01-02",
                "collections": ["示例文集"],
                "categories": ["自定义分类"],
                "comments": {"paragraph": False, "article": True},
                "customField": "保留",
            }
            (first / "index.md").write_text(
                source(
                    {
                        **common,
                        "title": "第一篇",
                        "date": "2023-03-04",
                        "slug": "first",
                        "weight": 10,
                        "articleId": "article-1111111111111111",
                        "articleRevision": "article-1111111111111111@old",
                    }
                ),
                encoding="utf-8",
            )
            (second / "index.md").write_text(
                source(
                    {
                        **common,
                        "title": "第二篇",
                        "date": "2023-05-06",
                        "slug": "second",
                        "weight": 20,
                        "articleId": "article-2222222222222222",
                        "articleRevision": "article-2222222222222222@old",
                    }
                ),
                encoding="utf-8",
            )

            document = Document()
            document.add_heading("第一篇", level=1)
            document.add_paragraph("第一篇的新正文。")
            document.add_heading("第二篇", level=1)
            document.add_paragraph("第二篇的新正文。")
            word = project / "示例文集.docx"
            document.save(word)

            completed = subprocess.run(
                [
                    sys.executable,
                    str(IMPORTER),
                    str(word),
                    "--project-root",
                    str(project),
                    "--collection-title",
                    "示例文集",
                    "--slugs",
                    "first,second",
                    "--date",
                    "2026-08-02",
                    "--backup-dir",
                    str(project / "backups"),
                ],
                text=True,
                capture_output=True,
                timeout=60,
            )
            self.assertEqual(completed.returncode, 0, completed.stderr)

            first_data = front_matter(first / "index.md")
            second_data = front_matter(second / "index.md")
            collection_data = front_matter(collection / "_index.md")
            self.assertEqual(first_data["date"], "2023-03-04")
            self.assertEqual(second_data["date"], "2023-05-06")
            self.assertEqual(first_data["lastmod"], "2026-08-02")
            self.assertEqual(first_data["articleId"], "article-1111111111111111")
            self.assertEqual(first_data["categories"], ["自定义分类"])
            self.assertEqual(first_data["comments"], {"paragraph": False, "article": True})
            self.assertEqual(first_data["customField"], "保留")
            self.assertNotEqual(first_data["articleRevision"], "article-1111111111111111@old")
            self.assertEqual(collection_data["date"], "2024-01-01")
            self.assertEqual(collection_data["endDate"], "2024-01-01")
            self.assertEqual(collection_data["lastmod"], "2026-08-02")
            self.assertEqual(collection_data["customCollectionField"], "保留")


if __name__ == "__main__":
    unittest.main()
