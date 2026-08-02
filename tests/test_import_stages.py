#!/usr/bin/env python3
"""DOCX 导入三阶段管线（parse/plan/commit）与机器接口的测试。"""

from __future__ import annotations

import copy
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

import yaml
from PIL import Image
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "importer"))

import import_stages  # noqa: E402

API = ROOT / "importer" / "import_api.py"


def front_matter(markdown: str) -> dict:
    _, raw, _ = markdown.split("---", 2)
    return yaml.safe_load(raw)


class ImportStageTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="writing-stages-test-")
        self.root = Path(self.temp.name)
        (self.root / "content").mkdir()
        self.source = self.root / "《测试文集·春日记》.docx"

    def tearDown(self):
        self.temp.cleanup()

    def basic_document(self, path=None):
        doc = Document()
        doc.add_paragraph("春日记", style="Title")
        doc.add_heading("第一章", level=1)
        doc.add_paragraph("这是一个保留中文标点的自然段。")
        doc.save(path or self.source)

    def make_options(self, **overrides):
        base = dict(
            title="春日记",
            slug="spring-notes",
            section="works",
            collections=["示例文集"],
            collection_slug="sample-works",
        )
        base.update(overrides)
        return import_stages.ImportOptions(**base)

    def tree_snapshot(self):
        return sorted(str(p.relative_to(self.root)) for p in self.root.rglob("*"))

    def temp_import_dirs(self):
        return sorted(self.root.rglob("*.import-*"))

    def run_api(self, *args: str):
        command = [sys.executable, str(API), "--project-root", str(self.root), *args]
        # 机器接口测试必须与调用测试的终端形态无关。显式关闭标准输入，
        # 才能真实覆盖“无 TTY 且未传 --yes 时安全拒绝”的分支。
        return subprocess.run(
            command,
            text=True,
            stdin=subprocess.DEVNULL,
            capture_output=True,
            timeout=60,
        )

    def test_01解析阶段不写任何目录(self):
        self.basic_document()
        before = self.tree_snapshot()
        result = import_stages.parse_docx(self.source)
        self.assertEqual(self.tree_snapshot(), before)
        self.assertEqual(result.source.filename, self.source.name)
        self.assertEqual(len(result.source.sha256), 64)
        self.assertEqual(result.word_title, "春日记")
        self.assertIn("自然段", result.body_markdown)
        self.assertEqual(result.heading_count, 1)
        self.assertGreaterEqual(result.paragraph_count, 2)
        kinds = {block.kind for block in result.blocks}
        self.assertIn("heading-1", kinds)
        self.assertIn("paragraph", kinds)

    def test_02计划阶段不写盘且同输入输出完全一致(self):
        self.basic_document()
        before = self.tree_snapshot()
        first = import_stages.plan_import(import_stages.parse_docx(self.source), self.make_options(), self.root)
        second = import_stages.plan_import(import_stages.parse_docx(self.source), self.make_options(), self.root)
        self.assertEqual(self.tree_snapshot(), before)
        dump_first = json.dumps(first.to_json_dict(), ensure_ascii=False, indent=2)
        dump_second = json.dumps(second.to_json_dict(), ensure_ascii=False, indent=2)
        self.assertEqual(dump_first, dump_second)
        data = first.to_json_dict()
        self.assertTrue(data["frontMatter"]["articleId"].startswith("article-"))
        self.assertIn("<!-- paragraph-id:p-", data["markdown"])
        self.assertEqual(data["target"], "content/works/sample-works/spring-notes")
        self.assertIn("content/works/sample-works/spring-notes/index.md", data["proposedFiles"])
        self.assertIn("content/works/sample-works/_index.md", data["proposedFiles"])
        self.assertNotIn(str(self.root), dump_first)

    def test_03冲突检测目标存在slug重复与articleId重复(self):
        self.basic_document()
        occupied = self.root / "content" / "works" / "sample-works" / "spring-notes"
        occupied.mkdir(parents=True)
        plan = import_stages.plan_import(import_stages.parse_docx(self.source), self.make_options(), self.root)
        self.assertIn("target-exists", {conflict.code for conflict in plan.conflicts})

        other = self.root / "content" / "essays" / "other"
        other.mkdir(parents=True)
        (other / "index.md").write_text(
            '---\ntitle: "别的文章"\nslug: "spring-notes"\narticleId: "article-0000000000000000"\n---\n正文\n',
            encoding="utf-8",
        )
        plan = import_stages.plan_import(
            import_stages.parse_docx(self.source),
            self.make_options(article_id="article-0000000000000000"),
            self.root,
        )
        codes = {conflict.code for conflict in plan.conflicts}
        self.assertIn("slug-duplicate", codes)
        self.assertIn("article-id-duplicate", codes)

    def test_04非法slug与路径穿越是致命错误而非冲突(self):
        self.basic_document()
        with self.assertRaises(import_stages.ValidationFailure):
            import_stages.plan_import(
                import_stages.parse_docx(self.source), self.make_options(slug="../escape"), self.root
            )
        with self.assertRaises(import_stages.ValidationFailure):
            import_stages.plan_import(
                import_stages.parse_docx(self.source), self.make_options(slug="Bad Slug"), self.root
            )

    def test_05提交成功文件完整且图片字节一致(self):
        image = self.root / "source.png"
        Image.new("RGB", (640, 480), "#9a4d3a").save(image)
        doc = Document()
        doc.add_paragraph("含图片的正文。")
        doc.add_picture(str(image))
        doc.save(self.source)
        parsed = import_stages.parse_docx(self.source)
        plan = import_stages.plan_import(parsed, self.make_options(), self.root)
        result = import_stages.commit_import(plan, parsed, self.root)
        bundle = Path(result.output)
        markdown = (bundle / "index.md").read_text(encoding="utf-8")
        data = front_matter(markdown)
        self.assertEqual(data["title"], "春日记")
        self.assertEqual(data["articleId"], plan.front_matter["articleId"])
        self.assertEqual(data["articleRevision"], plan.front_matter["articleRevision"])
        self.assertIn("<!-- paragraph-id:", markdown)
        original = bundle / "images" / "original" / "image-001.png"
        self.assertEqual(original.read_bytes(), image.read_bytes())
        webp = bundle / "images" / "image-001.webp"
        self.assertTrue(webp.is_file())
        sizes = {asset.target_name: asset.size for asset in plan.assets}
        self.assertEqual(webp.stat().st_size, sizes[f"{plan.target}/images/image-001.webp"])
        self.assertEqual(original.stat().st_size, sizes[f"{plan.target}/images/original/image-001.png"])
        collection = self.root / "content" / "works" / "sample-works" / "_index.md"
        self.assertTrue(collection.is_file())
        # 成功后不残留临时目录
        self.assertEqual(self.temp_import_dirs(), [])

    def test_06提交遇到冲突拒绝且不留半成品(self):
        self.basic_document()
        occupied = self.root / "content" / "works" / "sample-works" / "spring-notes"
        occupied.mkdir(parents=True)
        (occupied / "index.md").write_text("旧稿\n", encoding="utf-8")
        parsed = import_stages.parse_docx(self.source)
        plan = import_stages.plan_import(parsed, self.make_options(), self.root)
        with self.assertRaises(import_stages.ConflictFailure):
            import_stages.commit_import(plan, parsed, self.root)
        self.assertEqual((occupied / "index.md").read_text(encoding="utf-8"), "旧稿\n")
        self.assertEqual(self.temp_import_dirs(), [])

    def test_07提交写入中途失败时回滚并清理临时目录(self):
        self.basic_document()
        parsed = import_stages.parse_docx(self.source)
        plan = import_stages.plan_import(parsed, self.make_options(), self.root)
        with mock.patch.object(import_stages, "os_replace", side_effect=OSError("模拟写入失败")):
            with self.assertRaises(OSError):
                import_stages.commit_import(plan, parsed, self.root)
        target = self.root / "content" / "works" / "sample-works" / "spring-notes"
        self.assertFalse(target.exists())
        self.assertFalse((self.root / "content" / "works" / "sample-works" / "_index.md").exists())
        self.assertEqual(self.temp_import_dirs(), [])

    def test_08伪装docx被拒绝(self):
        fake = self.root / "fake.docx"
        fake.write_text("这不是一个docx文件", encoding="utf-8")
        with self.assertRaises(import_stages.SourceFailure):
            import_stages.parse_docx(fake)
        completed = self.run_api("inspect", str(fake), "--json")
        self.assertEqual(completed.returncode, 2)
        payload = json.loads(completed.stdout)
        self.assertFalse(payload["ok"])
        self.assertEqual(payload["error"]["code"], "invalid-source")

    def test_09空文档被拒绝(self):
        Document().save(self.source)
        parsed = import_stages.parse_docx(self.source)
        with self.assertRaises(import_stages.ValidationFailure):
            import_stages.plan_import(parsed, self.make_options(), self.root)

    def test_10中文空格与特殊字符文件名可以完整走通(self):
        special = self.root / "《测试文集·春日 记（终稿） v2》.docx"
        self.basic_document(special)
        completed = self.run_api(
            "plan",
            str(special),
            "--json",
            "--title",
            "春日记",
            "--slug",
            "spring-notes",
            "--section",
            "works",
            "--collection",
            "示例文集",
            "--collection-slug",
            "sample-works",
            "--save",
            str(self.root / "plan.json"),
        )
        self.assertEqual(completed.returncode, 0, completed.stderr)
        payload = json.loads(completed.stdout)
        self.assertTrue(payload["ok"])
        self.assertEqual(payload["plan"]["document"]["suggestedTitle"], "春日 记（终稿） v2")
        self.assertNotIn(str(self.root), completed.stdout)
        committed = self.run_api(
            "commit", str(self.root / "plan.json"), "--source", str(special), "--yes", "--json"
        )
        self.assertEqual(committed.returncode, 0, committed.stderr)
        result = json.loads(committed.stdout)
        self.assertTrue(result["ok"])
        self.assertTrue(
            (self.root / "content" / "works" / "sample-works" / "spring-notes" / "index.md").is_file()
        )

    def test_11同一关系ID的重复图片只收集一次(self):
        image = self.root / "dup.png"
        Image.new("RGB", (320, 200), "#3a4d9a").save(image)
        doc = Document()
        doc.add_picture(str(image))
        picture_paragraph = doc.paragraphs[0]
        extra = doc.add_paragraph()
        for child in picture_paragraph._p.iterchildren():
            extra._p.append(copy.deepcopy(child))
        doc.save(self.source)
        parsed = import_stages.parse_docx(self.source)
        self.assertEqual(len(parsed.images), 1)
        self.assertEqual(parsed.body_markdown.count("images/image-001.webp"), 2)

    def test_12未知样式与脚注产生警告但不改渲染(self):
        doc = Document()
        doc.styles.add_style("花样式", WD_STYLE_TYPE.PARAGRAPH)
        doc.add_paragraph("使用自定义样式的段落。", style="花样式")
        doc.save(self.source)
        footnotes_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:footnote w:type="separator" w:id="-1"><w:p/></w:footnote>'
            '<w:footnote w:type="continuationSeparator" w:id="0"><w:p/></w:footnote>'
            '<w:footnote w:id="1"><w:p><w:r><w:t>脚注内容</w:t></w:r></w:p></w:footnote>'
            "</w:footnotes>"
        )
        with zipfile.ZipFile(self.source, "a") as archive:
            archive.writestr("word/footnotes.xml", footnotes_xml)
        parsed = import_stages.parse_docx(self.source)
        codes = {warning.code for warning in parsed.warnings}
        self.assertIn("unknown-style", codes)
        self.assertIn("footnote-unsupported", codes)
        unknown = next(warning for warning in parsed.warnings if warning.code == "unknown-style")
        self.assertIn("花样式", unknown.message)
        self.assertIn("第1段", unknown.location)
        self.assertEqual(parsed.footnote_count, 1)
        self.assertIn("使用自定义样式的段落。", parsed.body_markdown)

    def test_13机器接口commit校验sha冲突与确认(self):
        self.basic_document()
        planned = self.run_api(
            "plan",
            str(self.source),
            "--json",
            "--title",
            "春日记",
            "--slug",
            "spring-notes",
            "--section",
            "works",
            "--collection",
            "示例文集",
            "--collection-slug",
            "sample-works",
            "--save",
            str(self.root / "plan.json"),
        )
        self.assertEqual(planned.returncode, 0, planned.stderr)
        # 无tty且没有--yes：拒绝写入
        refused = self.run_api("commit", str(self.root / "plan.json"), "--source", str(self.source))
        self.assertEqual(refused.returncode, 3)
        self.assertEqual(json.loads(refused.stdout)["error"]["code"], "confirmation-required")
        # sha256不匹配：源文件内容被改动后拒绝
        changed = Document()
        changed.add_paragraph("这是被改动后的另一份文档。")
        changed.save(self.source)
        mismatched = self.run_api(
            "commit", str(self.root / "plan.json"), "--source", str(self.source), "--yes", "--json"
        )
        self.assertEqual(mismatched.returncode, 3)
        self.assertEqual(json.loads(mismatched.stdout)["error"]["code"], "validation-failed")

    def test_14机器接口commit遇到计划内冲突退出码为4(self):
        self.basic_document()
        occupied = self.root / "content" / "works" / "sample-works" / "spring-notes"
        occupied.mkdir(parents=True)
        planned = self.run_api(
            "plan",
            str(self.source),
            "--json",
            "--title",
            "春日记",
            "--slug",
            "spring-notes",
            "--section",
            "works",
            "--collection",
            "示例文集",
            "--collection-slug",
            "sample-works",
            "--save",
            str(self.root / "plan.json"),
        )
        self.assertEqual(planned.returncode, 0, planned.stderr)
        payload = json.loads(planned.stdout)
        self.assertIn("target-exists", {item["code"] for item in payload["plan"]["conflicts"]})
        committed = self.run_api(
            "commit", str(self.root / "plan.json"), "--source", str(self.source), "--yes", "--json"
        )
        self.assertEqual(committed.returncode, 4)
        self.assertEqual(json.loads(committed.stdout)["error"]["code"], "conflict")

    def test_15机器接口plan非法slug退出码为3(self):
        self.basic_document()
        completed = self.run_api(
            "plan",
            str(self.source),
            "--json",
            "--title",
            "春日记",
            "--slug",
            "../escape",
            "--section",
            "essays",
        )
        self.assertEqual(completed.returncode, 3)
        payload = json.loads(completed.stdout)
        self.assertFalse(payload["ok"])
        self.assertEqual(payload["error"]["code"], "validation-failed")

    def test_16inspect输出建议与统计(self):
        self.basic_document()
        completed = self.run_api("inspect", str(self.source), "--json")
        self.assertEqual(completed.returncode, 0, completed.stderr)
        payload = json.loads(completed.stdout)
        self.assertTrue(payload["ok"])
        self.assertEqual(payload["document"]["suggestedTitle"], "春日记")
        self.assertEqual(payload["document"]["suggestedSlug"], "chun-ri-ji")
        self.assertEqual(payload["document"]["suggestedSection"], "works")
        self.assertGreater(payload["document"]["wordCount"], 0)
        self.assertEqual(len(payload["source"]["sha256"]), 64)


if __name__ == "__main__":
    unittest.main(verbosity=2)
