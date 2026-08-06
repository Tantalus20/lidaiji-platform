#!/usr/bin/env python3
"""作者工作台（studio/）的端到端测试：临时项目骨架 + 随机空闲端口。

覆盖：回环绑定、上传校验、CSRF 防线、content/ 不被提前写入、
inspect/plan/commit 全链路、冲突 409 且不覆盖、临时目录清理、
图片接口白名单、前端无浏览器弹窗、元数据路径字段拒绝。
"""

from __future__ import annotations

import http.client
import json
import shutil
import sys
import tempfile
import threading
import unittest
from pathlib import Path

import yaml
from PIL import Image
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "importer"))

import studio.server as studio_server  # noqa: E402

studio_server.StudioHandler.log_message = lambda *args: None  # 测试中静默访问日志

BOUNDARY = "----studio-test-boundary"


def make_docx(path: Path, with_image: bool = False) -> None:
    document = Document()
    document.add_paragraph("测试稿", style="Title")
    document.add_heading("第一章", level=1)
    document.add_paragraph("这是一段用于作者工作台测试的正文。")
    if with_image:
        image_path = path.parent / "fixture.png"
        Image.new("RGB", (320, 200), "#9a4d3a").save(image_path)
        document.add_picture(str(image_path))
    document.save(path)


def multipart_body(filename: str, data: bytes) -> tuple[bytes, str]:
    body = (
        f"--{BOUNDARY}\r\n".encode()
        + f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'.encode()
        + b"Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n"
        + data
        + b"\r\n"
        + f"--{BOUNDARY}--\r\n".encode()
    )
    return body, f"multipart/form-data; boundary={BOUNDARY}"


def front_matter(markdown: str) -> dict:
    _, raw, _ = markdown.split("---", 2)
    return yaml.safe_load(raw)


class StudioTestCase(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="studio-test-")
        self.root = Path(self.temp.name)
        (self.root / "content" / "works").mkdir(parents=True)
        (self.root / "content" / "essays").mkdir(parents=True)
        self.server = studio_server.create_server(self.root, port=0)
        self.server.state.auth_mode = "password"
        self.host, self.port = self.server.server_address[:2]
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.docx = self.root / "《测试文集·测试稿》.docx"
        make_docx(self.docx, with_image=True)

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=5)
        self.temp.cleanup()

    # -- 请求辅助 -----------------------------------------------------------

    def request(self, method: str, path: str, body: bytes = b"", headers: dict | None = None):
        connection = http.client.HTTPConnection(self.host, self.port, timeout=30)
        merged = {"Cookie": f"{studio_server.STUDIO_SESSION_COOKIE}={self.server.state.studio_session_token}"}
        merged.update(headers or {})
        connection.request(method, path, body, merged)
        response = connection.getresponse()
        payload = response.read()
        connection.close()
        return response.status, payload

    def request_without_automatic_host(
        self, method: str, path: str, headers: list[tuple[str, str]] | None = None, body: bytes = b""
    ):
        """发送不由http.client补写Host的请求，用于覆盖缺失和重复Host。"""
        connection = http.client.HTTPConnection(self.host, self.port, timeout=30)
        connection.putrequest(method, path, skip_host=True)
        for name, value in headers or []:
            connection.putheader(name, value)
        if body:
            connection.putheader("Content-Length", str(len(body)))
        connection.endheaders(body)
        response = connection.getresponse()
        payload = response.read()
        connection.close()
        return response.status, payload

    def post_json(self, path: str, payload: dict, headers: dict | None = None):
        merged = {
            "Content-Type": "application/json",
            "X-Studio-Request": "1",
            "X-Studio-CSRF": self.server.state.studio_csrf_token,
        }
        merged.update(headers or {})
        status, payload = self.request("POST", path, json.dumps(payload).encode("utf-8"), merged)
        return status, json.loads(payload)

    def upload(self, path: Path | None = None, filename: str | None = None, data: bytes | None = None):
        source = path or self.docx
        body, content_type = multipart_body(filename or source.name, data if data is not None else source.read_bytes())
        status, payload = self.request(
            "POST",
            "/api/import/inspect",
            body,
            {
                "Content-Type": content_type,
                "X-Studio-Request": "1",
                "X-Studio-CSRF": self.server.state.studio_csrf_token,
            },
        )
        return status, json.loads(payload)

    def plan(self, token: str, metadata: dict):
        return self.post_json("/api/import/plan", {"token": token, "metadata": metadata})

    def basic_metadata(self, **overrides):
        metadata = {"title": "测试稿", "slug": "studio-test-article", "section": "essays"}
        metadata.update(overrides)
        return metadata

    def content_snapshot(self):
        return sorted(str(item.relative_to(self.root)) for item in (self.root / "content").rglob("*"))

    # -- 测试 ---------------------------------------------------------------

    def test_01绑定地址必须是回环(self):
        self.assertEqual(self.host, "127.0.0.1")
        with self.assertRaises(ValueError):
            studio_server.create_server(self.root, host="0.0.0.0", port=0)
        with self.assertRaises(ValueError):
            studio_server.create_server(self.root, host="192.168.1.10", port=0)

    def test_02非docx上传被拒(self):
        status, payload = self.upload(filename="notes.txt")
        self.assertEqual(status, 400)
        self.assertEqual(payload["error"]["code"], "invalid-source")
        status, payload = self.upload(filename="fake.docx", data="这不是docx".encode("utf-8"))
        self.assertEqual(status, 400)
        self.assertEqual(payload["error"]["code"], "invalid-source")

    def test_03缺自定义头与错误Origin被403(self):
        body, content_type = multipart_body(self.docx.name, self.docx.read_bytes())
        status, _ = self.request("POST", "/api/import/inspect", body, {"Content-Type": content_type})
        self.assertEqual(status, 403)
        status, payload = self.post_json(
            "/api/import/abort", {"token": "0" * 32}, {"Origin": "http://evil.example.com"}
        )
        self.assertEqual(status, 403)
        self.assertEqual(payload["error"]["code"], "forbidden")
        # 回环 Origin 放行（继续走到业务校验而不是 403）
        status, _ = self.post_json(
            "/api/import/abort", {"token": "0" * 32}, {"Origin": f"http://127.0.0.1:{self.port}"}
        )
        self.assertEqual(status, 404)  # token 不存在，说明已通过 CSRF 闸

    def test_04上传只落临时目录不动content(self):
        before = self.content_snapshot()
        status, payload = self.upload()
        self.assertEqual(status, 200, payload)
        self.assertEqual(self.content_snapshot(), before)
        session_dir = self.root / ".cache" / "studio" / payload["token"]
        self.assertTrue((session_dir / "source.docx").is_file())

    def test_05inspect返回报告字段(self):
        status, payload = self.upload()
        self.assertTrue(payload["ok"])
        report = payload["report"]
        self.assertEqual(len(report["source"]["sha256"]), 64)
        self.assertGreater(report["document"]["wordCount"], 0)
        self.assertEqual(report["document"]["imageCount"], 1)
        self.assertIn("suggestedTitle", report["document"])
        self.assertIsInstance(report["warnings"], list)
        self.assertIsInstance(report["collections"], list)

    def test_06plan默认草稿(self):
        _, payload = self.upload()
        status, planned = self.plan(payload["token"], self.basic_metadata())
        self.assertEqual(status, 200, planned)
        self.assertTrue(planned["plan"]["frontMatter"]["draft"])
        self.assertIn("previewHtml", planned)
        self.assertTrue(planned["assetUrls"])
        self.assertEqual(planned["plan"]["conflicts"], [])

    def test_07非法slug被拒(self):
        _, payload = self.upload()
        status, result = self.plan(payload["token"], self.basic_metadata(slug="../escape"))
        self.assertEqual(status, 400)
        self.assertEqual(result["error"]["code"], "validation-failed")
        status, result = self.plan(payload["token"], self.basic_metadata(slug="Bad Slug"))
        self.assertEqual(status, 400)
        self.assertEqual(result["error"]["code"], "validation-failed")

    def test_08commit成功生成草稿bundle(self):
        _, payload = self.upload()
        _, planned = self.plan(payload["token"], self.basic_metadata())
        status, committed = self.post_json("/api/import/commit", {"token": payload["token"]})
        self.assertEqual(status, 200, committed)
        target = committed["result"]["target"]
        self.assertEqual(target, "content/essays/studio-test-article")
        bundle = self.root / target
        markdown = (bundle / "index.md").read_text(encoding="utf-8")
        data = front_matter(markdown)
        self.assertTrue(data["draft"])
        self.assertEqual(data["title"], "测试稿")
        self.assertTrue((bundle / "images" / "image-001.webp").is_file())
        self.assertTrue((bundle / "images" / "original" / "image-001.png").is_file())
        self.assertEqual(planned["plan"]["proposedFiles"][0], f"{target}/index.md")

    def test_09重复slug冲突409且不覆盖(self):
        _, first = self.upload()
        self.plan(first["token"], self.basic_metadata())
        status, committed = self.post_json("/api/import/commit", {"token": first["token"]})
        self.assertEqual(status, 200)
        index = self.root / "content" / "essays" / "studio-test-article" / "index.md"
        original_bytes = index.read_bytes()
        _, second = self.upload()
        _, planned = self.plan(second["token"], self.basic_metadata())
        self.assertIn("target-exists", {item["code"] for item in planned["plan"]["conflicts"]})
        status, refused = self.post_json("/api/import/commit", {"token": second["token"]})
        self.assertEqual(status, 409)
        self.assertEqual(refused["error"]["code"], "conflict")
        self.assertEqual(index.read_bytes(), original_bytes)

    def test_10失败后无半成品(self):
        _, first = self.upload()
        self.plan(first["token"], self.basic_metadata())
        self.post_json("/api/import/commit", {"token": first["token"]})
        _, second = self.upload()
        self.plan(second["token"], self.basic_metadata())
        status, _ = self.post_json("/api/import/commit", {"token": second["token"]})
        self.assertEqual(status, 409)
        self.assertEqual(list(self.root.rglob("*.import-*")), [])

    def test_11图片接口拒绝路径穿越(self):
        _, payload = self.upload()
        token = payload["token"]
        status, _ = self.request("GET", f"/api/import/asset?token={token}&name=../x")
        self.assertEqual(status, 400)
        status, _ = self.request("GET", f"/api/import/asset?token={token}&name=image-001.webp")
        self.assertEqual(status, 200)
        status, _ = self.request("GET", f"/api/import/asset?token={'1' * 32}&name=image-001.webp")
        self.assertEqual(status, 404)

    def test_12前端源码不使用浏览器弹窗(self):
        static_dir = ROOT / "studio" / "static"
        for file in static_dir.iterdir():
            if not file.is_file():
                continue
            self.assertNotIn("al" + "ert(", file.read_text(encoding="utf-8"), f"{file.name} 含有弹窗调用")

    def test_12b页面引用的静态资源必须都在服务端白名单中(self):
        import re
        html = (ROOT / "studio" / "static" / "index.html").read_text(encoding="utf-8")
        references = re.findall(r'(?:src|href)="(/[^"#]+)"', html)
        references = [item for item in references if not item.startswith("/api/")]
        self.assertTrue(references, "index.html 未找到任何静态资源引用")
        for reference in references:
            self.assertIn(reference, studio_server.STATIC_FILES, f"页面引用 {reference} 不在服务端静态白名单")
            filename, _ = studio_server.STATIC_FILES[reference]
            self.assertTrue(
                (ROOT / "studio" / "static" / filename).is_file(),
                f"{reference} 对应的静态文件不存在：{filename}",
            )

    def test_12c静态响应带CSP与nosniff(self):
        connection = http.client.HTTPConnection(self.host, self.port, timeout=30)
        connection.request("GET", "/")
        response = connection.getresponse()
        response.read()
        csp = response.getheader("Content-Security-Policy", "")
        self.assertIn("default-src 'none'", csp)
        self.assertIn("script-src 'self'", csp)
        self.assertIn("object-src 'none'", csp)
        self.assertEqual(response.getheader("X-Content-Type-Options"), "nosniff")
        connection.close()

    def test_13关闭服务后会话目录被清理(self):
        _, payload = self.upload()
        self.assertTrue((self.root / ".cache" / "studio" / payload["token"]).is_dir())
        self.server.server_close()
        # 会话临时目录必须清理；持久状态（账本/候选/发布日志）保留
        self.assertFalse((self.root / ".cache" / "studio" / payload["token"]).exists())

    def test_14元数据路径字段与穿越值被拒(self):
        _, payload = self.upload()
        status, result = self.plan(
            payload["token"],
            self.basic_metadata(
                section="works",
                collections=["测试文集"],
                collectionSlug="../../etc",
            ),
        )
        self.assertEqual(status, 400)
        self.assertEqual(result["error"]["code"], "validation-failed")
        status, result = self.plan(
            payload["token"], self.basic_metadata(path="/etc/passwd", project_root="/tmp")
        )
        self.assertEqual(status, 400)
        self.assertEqual(result["error"]["code"], "validation-failed")
        self.assertEqual(self.content_snapshot(), ["content/essays", "content/works"])

    def test_15系统状态与预览接口形状(self):
        status, payload = self.request("GET", "/api/system/status")
        self.assertEqual(status, 200)
        body = json.loads(payload)
        self.assertTrue(body["ok"])
        self.assertFalse(body["preview"]["running"])
        self.assertEqual(body["preview"]["url"], "http://127.0.0.1:1313/")
        status, payload = self.post_json("/api/system/preview", {"action": "bogus"})
        self.assertEqual(status, 400)

    def test_16abort清理会话(self):
        _, payload = self.upload()
        token = payload["token"]
        status, aborted = self.post_json("/api/import/abort", {"token": token})
        self.assertEqual(status, 200)
        self.assertTrue(aborted["ok"])
        self.assertFalse((self.root / ".cache" / "studio" / token).exists())
        status, _ = self.post_json("/api/import/commit", {"token": token})
        self.assertEqual(status, 404)

    def test_17合法回环Host均可访问(self):
        for authority in (
            f"127.0.0.1:{self.port}",
            f"localhost:{self.port}",
            f"[::1]:{self.port}",
        ):
            with self.subTest(authority=authority):
                status, payload = self.request("GET", "/", headers={"Host": authority})
                self.assertEqual(status, 200)
                self.assertIn(b"<!doctype html>", payload.lower())

    def test_18非回环Host不能读取页面或API(self):
        cases = (
            ("/", f"rebind.example:{self.port}"),
            ("/api/articles", f"rebind.example:{self.port}"),
            ("/api/git/diff?path=content/works/example/index.md", f"evil.example:{self.port}"),
            ("/api/media", f"evil.example:{self.port}"),
            ("/api/articles", f"localhost.evil.example:{self.port}"),
            ("/api/articles", f"127.0.0.1.evil.example:{self.port}"),
            ("/api/articles", f"localhost@evil.example:{self.port}"),
            ("/api/articles", f"127.0.0.1:{self.port}.evil.example"),
        )
        for path, authority in cases:
            with self.subTest(path=path, authority=authority):
                status, payload = self.request("GET", path, headers={"Host": authority})
                self.assertEqual(status, 403)
                self.assertNotIn("测试稿".encode(), payload)

    def test_19非回环Host不能执行POST(self):
        body = json.dumps({"token": "0" * 32}).encode("utf-8")
        status, _ = self.request(
            "POST",
            "/api/import/abort",
            body,
            {
                "Host": f"rebind.example:{self.port}",
                "Content-Type": "application/json",
                "X-Studio-Request": "1",
            },
        )
        self.assertEqual(status, 403)

    def test_20缺失重复和畸形Host均被拒绝(self):
        status, _ = self.request_without_automatic_host("GET", "/")
        self.assertEqual(status, 403)
        cookie = f"{studio_server.STUDIO_SESSION_COOKIE}={self.server.state.studio_session_token}"
        status, _ = self.request_without_automatic_host(
            "GET",
            "/api/articles",
            [
                ("Host", f"127.0.0.1:{self.port}"),
                ("Host", f"evil.example:{self.port}"),
                ("Cookie", cookie),
            ],
        )
        self.assertEqual(status, 403)
        status, _ = self.request("GET", "/", headers={"Host": "[::1"})
        self.assertEqual(status, 403)

    def test_21无有效Studio会话不能读取API(self):
        status, payload = self.request("GET", "/api/articles", headers={"Cookie": ""})
        self.assertEqual(status, 403)
        self.assertNotIn("测试稿".encode(), payload)

    def test_22首页初始化仅存在内存的会话Cookie(self):
        connection = http.client.HTTPConnection(self.host, self.port, timeout=30)
        connection.request("GET", "/")
        response = connection.getresponse()
        body = response.read()
        cookie = response.getheader("Set-Cookie") or ""
        connection.close()
        self.assertEqual(response.status, 200)
        self.assertIn(f"{studio_server.STUDIO_SESSION_COOKIE}=", cookie)
        self.assertIn("HttpOnly", cookie)
        self.assertIn("SameSite=Strict", cookie)
        self.assertNotIn(self.server.state.studio_session_token.encode(), body)

    def test_23Studio重启后旧会话失效(self):
        old_token = self.server.state.studio_session_token
        replacement = studio_server.create_server(self.root, port=0)
        host, port = replacement.server_address[:2]
        thread = threading.Thread(target=replacement.serve_forever, daemon=True)
        thread.start()
        try:
            connection = http.client.HTTPConnection(host, port, timeout=30)
            connection.request(
                "GET",
                "/api/articles",
                headers={"Cookie": f"{studio_server.STUDIO_SESSION_COOKIE}={old_token}"},
            )
            response = connection.getresponse()
            response.read()
            connection.close()
            self.assertEqual(response.status, 403)
            self.assertNotEqual(old_token, replacement.state.studio_session_token)
        finally:
            replacement.shutdown()
            replacement.server_close()
            thread.join(timeout=5)


if __name__ == "__main__":
    unittest.main(verbosity=2)
